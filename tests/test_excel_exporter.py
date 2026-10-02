"""Tests for app/exporters/excel_exporter.py — no live requests, uses tmp_path."""
import pytest
from datetime import datetime
from pathlib import Path

import openpyxl

from app.exporters.excel_exporter import export_excel
from app.models.lead import Lead


def _sample_leads():
    return [
        Lead(
            company_name="Earth Science Ayurveda",
            contact_person="Ms Jyoti Karala",
            designation="Marketing Manager",
            phone_number=None,
            email=None,
            address="Plot No.20, Block-A, Faridabad, Haryana, 121003",
            city="Faridabad",
            state="Haryana",
            gst_no="06AOGPR9376E2ZH",
            product_name="Rose Water 100ml",
            product_url="https://www.tradeindia.com/products/rose-water-c9649318.html",
            source="TradeIndia",
            phone_status="HIDDEN",
            email_status="MISSING",
            scraped_at=datetime(2026, 10, 2, 11, 0, 0),
        ),
        Lead(
            company_name="FALCON",
            contact_person="Mr Jaikishan L",
            designation="CEO",
            phone_number="+919876543210",
            email="contact@gmail.com",
            address="Bengaluru, Karnataka",
            city="Bengaluru",
            state="Karnataka",
            gst_no="29AEJPJ3915J1Z8",
            product_name="Rose Floral Water",
            product_url="https://www.tradeindia.com/products/rose-floral-water-c5735615.html",
            source="TradeIndia",
            phone_status="VALID",
            email_status="VALID",
            scraped_at=datetime(2026, 10, 2, 11, 0, 0),
        ),
    ]


class TestExcelOutput:
    def test_creates_file(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        result = export_excel(_sample_leads(), output_path=out)
        assert result == out
        assert out.exists()

    def test_correct_sheet_names(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        assert "Leads" in wb.sheetnames
        assert "Run Summary" in wb.sheetnames

    def test_row_count(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        # 1 header + 2 data rows
        assert ws.max_row == 3

    def test_column_count(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        assert ws.max_column == 15

    def test_header_values(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        headers = [ws.cell(1, c).value for c in range(1, 16)]
        assert "Company Name"   in headers
        assert "Phone Number"   in headers
        assert "GST No"         in headers
        assert "Designation"    in headers
        assert "Phone Status"   in headers
        assert "Email Status"   in headers
        assert "Product URL"    in headers

    def test_company_name_in_data(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        companies = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
        assert "Earth Science Ayurveda" in companies
        assert "FALCON" in companies

    def test_phone_stored_as_text(self, tmp_path):
        """Phone must be stored as text so +91 prefix isn't corrupted."""
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        # Find the Phone Number column
        headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        phone_col = headers["Phone Number"]
        # FALCON (row 3) has +919876543210
        cell = ws.cell(3, phone_col)
        assert str(cell.value) == "+919876543210"

    def test_product_url_has_hyperlink(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
        url_col = headers["Product URL"]
        cell = ws.cell(2, url_col)   # first data row
        assert cell.hyperlink is not None

    def test_frozen_header(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        assert ws.freeze_panes == "A2"

    def test_run_summary_totals(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel(_sample_leads(), output_path=out,
                     run_stats={"listings_found": 28, "leads_parsed": 28})
        wb = openpyxl.load_workbook(out)
        ws2 = wb["Run Summary"]
        # Collect all cell values
        all_vals = [ws2.cell(r, 2).value for r in range(1, ws2.max_row + 1)]
        assert 2 in all_vals    # Total leads = 2
        assert 28 in all_vals   # listings_found = 28

    def test_empty_leads_list(self, tmp_path):
        out = tmp_path / "test_leads.xlsx"
        export_excel([], output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        assert ws.max_row == 1   # header only


class TestExcelHardening:
    """Task 3: PermissionError fallback, Last Updated in summary, OneDrive path."""

    def test_last_updated_in_summary(self, tmp_path):
        """Run Summary must have a 'Last Updated' row."""
        out = tmp_path / "test.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws2 = wb["Run Summary"]
        all_labels = [str(ws2.cell(r, 1).value or "") for r in range(1, ws2.max_row + 1)]
        assert any("Last Updated" in lbl for lbl in all_labels), \
            f"'Last Updated' not found in summary labels: {all_labels}"

    def test_total_leads_in_summary(self, tmp_path):
        """Run Summary must show Total Leads count."""
        out = tmp_path / "test.xlsx"
        export_excel(_sample_leads(), output_path=out)
        wb = openpyxl.load_workbook(out)
        ws2 = wb["Run Summary"]
        all_labels = [str(ws2.cell(r, 1).value or "") for r in range(1, ws2.max_row + 1)]
        assert any("Total Leads" in lbl for lbl in all_labels), \
            f"'Total Leads' not found in summary labels: {all_labels}"

    def test_permission_error_writes_fallback(self, tmp_path):
        """If the target file is locked, a timestamped fallback is written."""
        from app.exporters.excel_exporter import _save_workbook
        import openpyxl as ox

        intended = tmp_path / "locked.xlsx"
        wb = ox.Workbook()

        # Simulate PermissionError by patching save on a temp copy
        original_save = ox.Workbook.save
        call_count = [0]

        def mock_save(self, filename):
            call_count[0] += 1
            if call_count[0] == 1:
                raise PermissionError("File is locked")
            original_save(self, filename)

        import unittest.mock as mock
        with mock.patch.object(ox.Workbook, "save", mock_save):
            actual = _save_workbook(wb, intended)

        # Fallback file must exist and be different from intended
        assert actual != intended
        assert actual.exists()
        assert actual.name.startswith("locked_")
        assert actual.suffix == ".xlsx"

    def test_permission_error_fallback_is_valid_xlsx(self, tmp_path):
        """The fallback file must be a readable .xlsx."""
        from app.exporters.excel_exporter import _save_workbook
        import openpyxl as ox
        import unittest.mock as mock

        intended = tmp_path / "locked.xlsx"
        wb = ox.Workbook()
        wb.active.title = "Test"
        wb.active["A1"] = "hello"

        original_save = ox.Workbook.save
        call_count = [0]

        def mock_save(self, filename):
            call_count[0] += 1
            if call_count[0] == 1:
                raise PermissionError("locked")
            original_save(self, filename)

        with mock.patch.object(ox.Workbook, "save", mock_save):
            actual = _save_workbook(wb, intended)

        wb2 = ox.load_workbook(actual)
        assert wb2["Test"]["A1"].value == "hello"

    def test_no_crash_when_primary_save_succeeds(self, tmp_path):
        """Normal save path still works end-to-end."""
        out = tmp_path / "normal.xlsx"
        result = export_excel(_sample_leads(), output_path=out)
        assert result == out
        assert out.exists()

    def test_onedrive_copy_called_when_configured(self, tmp_path):
        """If OUTPUT_FILE_ONEDRIVE is set, the file is copied there after save."""
        import shutil
        import unittest.mock as mock
        from app.exporters import excel_exporter as ee

        onedrive_dir = tmp_path / "onedrive"
        onedrive_dir.mkdir()
        onedrive_dest = onedrive_dir / "leads.xlsx"

        # Patch settings.onedrive_path to return our tmp dir path
        with mock.patch.object(
            type(ee.settings), "onedrive_path",
            new_callable=lambda: property(lambda self: onedrive_dest)
        ):
            out = tmp_path / "primary.xlsx"
            export_excel(_sample_leads(), output_path=out)

        assert onedrive_dest.exists(), "OneDrive copy was not created"
