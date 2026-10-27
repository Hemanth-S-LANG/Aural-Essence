"""
IndiaMART Lead Manager Scraper
==============================
Lead Manager → Contacts tab only.

For each contact:
  1. Click the contact card (left panel)
  2. Scrape the right panel:
       - Buyer phone (shown at top of chat header)
       - Product name, Quantity, Grade, Buyer Filled Details from the enquiry card
       - Chat messages
  3. Click "View More" (text link, top-right of chat)
  4. Scrape the popup modal:
       - Buyer name, phone, city, state, country
       - WhatsApp active status
       - Engagement: Requirements, Calls, Replies
       - Address
  5. Close modal (× button), move to next contact
  6. Stop when the list is exhausted (scroll reaches end)

Usage:
    python scrapeleads/scrape_leads.py

Output:
    scrapeleads/output/indiamart_leads.xlsx
    scrapeleads/output/indiamart_leads.csv
"""
import json, time, random, csv, re
from datetime import datetime
from pathlib import Path
from playwright.sync_api import sync_playwright

# ── Paths ─────────────────────────────────────────────────────────────────────
BASE_DIR    = Path(__file__).parent
OUT_DIR     = BASE_DIR / "output";  OUT_DIR.mkdir(exist_ok=True)
PROFILE_DIR = Path.home() / "IndiaMART_Profile";  PROFILE_DIR.mkdir(exist_ok=True)
OUT_XLSX    = OUT_DIR / "indiamart_leads.xlsx"
OUT_CSV     = OUT_DIR / "indiamart_leads.csv"
CKPT        = OUT_DIR / "checkpoint.json"
TARGET_URL  = "https://seller.indiamart.com/messagecentre"

# Seller's own numbers — filter these out so they don't appear as buyer phone
SELLER_PHONES = {"6361499510", "8047768879", "08047768879"}

# ── Checkpoint ────────────────────────────────────────────────────────────────
def load_ckpt():
    if CKPT.exists():
        try: return json.loads(CKPT.read_text(encoding="utf-8"))
        except: return []
    return []

def save_ckpt(leads):
    CKPT.write_text(json.dumps(leads, ensure_ascii=False, indent=2), encoding="utf-8")

# ── JS: ReactVirtualized helpers ──────────────────────────────────────────────

JS_LIST_INFO = r"""
() => {
    var list = document.querySelector('.ReactVirtualized__List');
    if (!list) return {ok: false};
    var inner = list.querySelector('.ReactVirtualized__Grid__innerScrollContainer');
    if (!inner) return {ok: false};
    var firstRow = inner.querySelector('div.row[role="gridcell"]');
    var rowH = firstRow ? (parseFloat(firstRow.style.height) || 120) : 120;
    var total = Math.round(list.scrollHeight / rowH);
    return {ok: true, total: total, rowH: rowH,
            scrollH: list.scrollHeight, clientH: list.clientHeight};
}
"""

JS_SCROLL_TO = r"""
(idx) => {
    var list = document.querySelector('.ReactVirtualized__List');
    if (!list) return false;
    var inner = list.querySelector('.ReactVirtualized__Grid__innerScrollContainer');
    var firstRow = inner ? inner.querySelector('div.row[role="gridcell"]') : null;
    var rowH = firstRow ? (parseFloat(firstRow.style.height) || 120) : 120;
    list.scrollTop = idx * rowH;
    return true;
}
"""

# Returns the rendered card at absolute position idx, or null if not in DOM yet
JS_GET_CARD = r"""
(idx) => {
    var inner = document.querySelector('.ReactVirtualized__Grid__innerScrollContainer');
    if (!inner) return null;
    var rows = inner.querySelectorAll('div.row[role="gridcell"]');
    var rowH = null;
    for (var i = 0; i < rows.length; i++) {
        if (!rowH && rows[i].style.height) rowH = parseFloat(rows[i].style.height) || 120;
        var top = parseFloat(rows[i].style.top) || 0;
        if (rowH && Math.round(top / rowH) === idx) {
            var el = rows[i].querySelector('[id^="contact-"]') || rows[i].firstElementChild;
            if (!el) return null;
            var lines = (el.innerText || '').split('\n').filter(function(l){return l.trim();});
            return {id: el.id, name: lines[0]||'', time: lines[1]||'',
                    location: lines[2]||'', preview: lines[3]||''};
        }
    }
    return null;
}
"""

JS_CLICK_CARD = r"""
(idx) => {
    var inner = document.querySelector('.ReactVirtualized__Grid__innerScrollContainer');
    if (!inner) return false;
    var rows = inner.querySelectorAll('div.row[role="gridcell"]');
    var rowH = null;
    for (var i = 0; i < rows.length; i++) {
        if (!rowH && rows[i].style.height) rowH = parseFloat(rows[i].style.height) || 120;
        var top = parseFloat(rows[i].style.top) || 0;
        if (rowH && Math.round(top / rowH) === idx) {
            var el = rows[i].querySelector('[id^="contact-"]') || rows[i].firstElementChild;
            if (el) { el.click(); return true; }
        }
    }
    return false;
}
"""

# ── JS: Right panel — phone + enquiry card + chat messages ───────────────────

JS_SCRAPE_PANEL = r"""
() => {
    var r = {phone:'', product:'', quantity:'', grade:'',
             buyer_filled:'', order_value:'', req_type:'', packaging:'',
             messages:[], enquiry_date:'', source:''};

    var body = document.body.innerText || '';

    // ── Phone: shown in the chat header as "06366568964" or "+91 636..."
    // It appears right after the buyer name line in the header area
    var phoneEl = document.querySelector(
        '.cntmsg_hdr_phn, [class*="cntmsg_hdr_phn"], ' +
        '[class*="phone_number"], [class*="phoneNum"], ' +
        '[class*="hdr_phone"], [class*="contact_phone"]'
    );
    if (phoneEl) r.phone = phoneEl.innerText.replace(/\D/g,'').trim();

    // Fallback: grab from href tel: links
    if (!r.phone) {
        var telLink = document.querySelector('[href^="tel:"]');
        if (telLink) r.phone = telLink.getAttribute('href').replace('tel:','').replace(/\D/g,'').trim();
    }

    // Fallback: regex on full body for 10-digit mobile NOT belonging to seller
    if (!r.phone) {
        var phones = body.match(/(?:\+?91[\s\-]?)?[6-9]\d{9}/g) || [];
        phones.forEach(function(p) {
            var d = p.replace(/\D/g,'');
            if (d.length === 12) d = d.slice(2);
            if (d.length === 10 && d !== '6361499510' && d !== '8047768879') {
                if (!r.phone) r.phone = d;
            }
        });
    }

    // ── Enquiry card fields (from the product card in the chat)
    var maps = {
        'product':      /^([^\n:]{3,60})\n\s*Quantity\s*:/m,
        'quantity':     /Quantity\s*:\s*([^\n]+)/i,
        'grade':        /Grade\s*(?:Standard)?\s*:\s*([^\n]+)/i,
        'buyer_filled': /Buyer\s+Filled\s+Details\s*:\s*([^\n]+)/i,
        'order_value':  /Probable\s+Order\s+Value\s*:\s*([^\n]+)/i,
        'req_type':     /Probable\s+Requirement\s+Type\s*:\s*([^\n]+)/i,
        'packaging':    /Packaging\s+Type\s*:\s*([^\n]+)/i,
    };
    for (var k in maps) {
        var m = body.match(maps[k]);
        if (m) r[k] = m[1].trim().substring(0, 200);
    }

    // ── Also try "Buyer Searched for..." as product fallback
    if (!r.product) {
        var bsm = body.match(/Buyer\s+(?:Searched|searched)\s+for\s+([^\n]+)/);
        if (bsm) r.product = bsm[0].trim().substring(0, 200);
    }

    // ── Chat messages from the right_side_msg container
    var msgEl = document.querySelector('.right_side_msg, [class*="right_side_msg"]');
    if (msgEl) {
        msgEl.querySelectorAll('[class*="msg_body"], [class*="msgBody"], p, .message-text, [class*="chat_msg"]').forEach(function(el) {
            var t = el.innerText.trim();
            if (t.length > 5 && t.length < 1000) r.messages.push(t);
        });
        // If no specific selectors matched, take full text
        if (r.messages.length === 0) {
            var t = msgEl.innerText.trim();
            if (t.length > 5) r.messages = [t.substring(0, 1000)];
        }
    }

    // ── Source (how enquiry arrived)
    var srcm = body.match(/Contact added through\s+([^\n]+)/i);
    if (srcm) r.source = srcm[0].trim();

    // ── Buyer name from chat header (appears before "last seen")
    var bnm = body.match(/([^\n]+)\n[^\n]*last seen/i);
    if (bnm) r.buyer_name_chat = bnm[1].trim();

    // ── Date visible in chat
    var dm = body.match(/\b(Today|Yesterday|\d{1,2}\s+\w+|\w+,\s+\d{1,2}\s+\w+)\b/);
    if (dm) r.enquiry_date = dm[0];

    return r;
}
"""

# ── JS: Click "View More" text link at top-right of chat ────────────────────

JS_CLICK_VIEWMORE = r"""
() => {
    // "View More" is a text link/button in the chat header area (top-right)
    // Photo shows it as plain text "View More" with blue styling
    var allEls = document.querySelectorAll(
        'span, a, button, div, label, p, [role="button"]'
    );
    for (var i = 0; i < allEls.length; i++) {
        var t = (allEls[i].innerText || '').trim();
        if (t === 'View More') {
            var rect = allEls[i].getBoundingClientRect();
            // Must be visible (width > 0, height > 0) and not inside the left panel
            if (rect.width > 0 && rect.height > 0 && rect.left > 400) {
                allEls[i].scrollIntoView({block:'center'});
                allEls[i].click();
                return {clicked: true, tag: allEls[i].tagName,
                        cls: (allEls[i].className||'').toString().substring(0,100),
                        left: Math.round(rect.left)};
            }
        }
    }
    return {clicked: false};
}
"""

# ── JS: Scrape the View More modal (photo 3) ─────────────────────────────────

JS_SCRAPE_MODAL = r"""
() => {
    // The modal from photo 3:
    // - Header: buyer avatar circle, name, "last seen", "This Month on platform"
    // - Phone number (e.g. 06366568964), location pin icon + "Bengaluru, Karnataka, India"
    // - "Chanchal is active on WhatsApp"
    // - Engagement section: Requirements (0), Calls (0), Replies (0)
    // - Contact Details section: ADDRESS field
    // - × close button top-right

    var r = {ok: false};

    // Find the modal — it's a fixed/absolute overlay with high z-index
    // From photo: it has a white card with rounded corners, × close button
    var modal = null;

    // Try role="dialog" first
    var dialogs = document.querySelectorAll('[role="dialog"], [role="alertdialog"]');
    for (var d = 0; d < dialogs.length; d++) {
        var rect = dialogs[d].getBoundingClientRect();
        if (rect.width > 200 && rect.height > 200) {
            modal = dialogs[d]; break;
        }
    }

    // Fallback: find highest z-index visible overlay
    if (!modal) {
        var best = null, bestZ = 0;
        document.querySelectorAll('div').forEach(function(el) {
            try {
                var s = window.getComputedStyle(el);
                var z = parseInt(s.zIndex) || 0;
                var pos = s.position;
                if ((pos === 'fixed' || pos === 'absolute') && z > 50) {
                    var rc = el.getBoundingClientRect();
                    if (rc.width > 300 && rc.height > 300 && z > bestZ) {
                        bestZ = z; best = el;
                    }
                }
            } catch(e) {}
        });
        if (best) modal = best;
    }

    if (!modal) {
        // Dump diagnostic for debugging
        var diag = [];
        document.querySelectorAll('div').forEach(function(el) {
            try {
                var s = window.getComputedStyle(el);
                var z = parseInt(s.zIndex)||0;
                if (z > 10) {
                    var rc = el.getBoundingClientRect();
                    if (rc.width > 100 && rc.height > 100) {
                        diag.push({z:z, pos:s.position, w:Math.round(rc.width),
                                   h:Math.round(rc.height), cls:(el.className||'').toString().substring(0,100)});
                    }
                }
            } catch(e) {}
        });
        diag.sort(function(a,b){return b.z-a.z;});
        return {ok: false, diag: diag.slice(0,8)};
    }

    r.ok = true;
    r._modal_cls = (modal.className||'').toString().substring(0,150);
    var txt = modal.innerText || '';
    r._modal_text = txt.substring(0, 3000);

    // ── Buyer name: first non-empty line in the modal
    var lines = txt.split('\n').map(function(l){return l.trim();}).filter(function(l){return l.length > 0;});
    r.buyer_name = lines[0] || '';

    // ── "This Month on platform"
    r.platform_period = '';
    var pm = txt.match(/(This Month|This Week|This Year|Last \d+ days)[^\n]*/i);
    if (pm) r.platform_period = pm[0].trim();

    // ── Phone: 10-digit number shown prominently
    r.phone = '';
    var phonEl = modal.querySelector('[href^="tel:"]');
    if (phonEl) {
        r.phone = phonEl.getAttribute('href').replace('tel:','').replace(/\D/g,'').trim();
    }
    if (!r.phone) {
        var pm2 = txt.match(/0?[6-9]\d{9}/);
        if (pm2) r.phone = pm2[0].replace(/\D/g,'').trim();
    }

    // ── Location (city, state, country line)
    r.location = '';
    var locEl = modal.querySelector('[class*="location"], [class*="city"], [class*="address"]');
    if (locEl) {
        r.location = locEl.innerText.trim();
    }
    if (!r.location) {
        // Look for the pin icon sibling text — typically "City, State, Country"
        var lm = txt.match(/([A-Z][a-zA-Z\s]+,\s*[A-Z][a-zA-Z\s]+,\s*India)/);
        if (lm) r.location = lm[1].trim();
    }

    // ── WhatsApp active
    r.whatsapp = '';
    var wam = txt.match(/([^\n]+is active on WhatsApp[^\n]*)/i);
    if (wam) r.whatsapp = wam[1].trim();

    // ── Engagement stats: Requirements, Calls, Replies
    r.req_count = ''; r.calls_count = ''; r.replies_count = '';
    var engM = txt.match(/(\d+)\s*\n?\s*REQUIREM(?:ENT)?S?/i);
    if (engM) r.req_count = engM[1];
    var callM = txt.match(/(\d+)\s*\n?\s*CALLS/i);
    if (callM) r.calls_count = callM[1];
    var repM = txt.match(/(\d+)\s*\n?\s*REPLIES/i);
    if (repM) r.replies_count = repM[1];

    // ── Address (Contact Details section)
    r.address = '';
    var addrEl = modal.querySelector('[class*="address"], [class*="Address"]');
    if (addrEl) r.address = addrEl.innerText.trim();
    if (!r.address && r.location) r.address = r.location;

    // ── Email
    r.email = '';
    var emlEl = modal.querySelector('[href^="mailto:"]');
    if (emlEl) r.email = emlEl.getAttribute('href').replace('mailto:','').trim();
    if (!r.email) {
        var em = txt.match(/[\w.+\-]+@[\w\-]+\.[\w.\-]+/);
        if (em) r.email = em[0];
    }

    return r;
}
"""

# ── JS: Close the modal ───────────────────────────────────────────────────────

JS_CLOSE_MODAL = r"""
() => {
    // Find × close button — it's top-right of the modal
    var closers = document.querySelectorAll(
        '[aria-label="Close"], [aria-label="close"], ' +
        '.ant-modal-close, [class*="closeIcon"], [class*="close_icon"], ' +
        '[class*="closeBtn"], [class*="close_btn"], [class*="modal_close"], ' +
        '[class*="modalClose"]'
    );
    for (var i = 0; i < closers.length; i++) {
        var rc = closers[i].getBoundingClientRect();
        if (rc.width > 0 && rc.height > 0) {
            closers[i].click();
            return {closed: true, method: 'button'};
        }
    }
    // Try finding × character as text
    var allBtns = document.querySelectorAll('button, span, div, [role="button"]');
    for (var j = 0; j < allBtns.length; j++) {
        var t = (allBtns[j].innerText || '').trim();
        if (t === '×' || t === '✕' || t === 'x' || t === 'X') {
            var rc2 = allBtns[j].getBoundingClientRect();
            if (rc2.width > 0 && rc2.height > 0) {
                allBtns[j].click();
                return {closed: true, method: 'x-char'};
            }
        }
    }
    return {closed: false};
}
"""


# ── Helpers ───────────────────────────────────────────────────────────────────

def clean_phone(raw: str) -> str:
    digits = re.sub(r"[^\d]", "", raw or "")
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) == 10 and digits[0] in "6789" and digits not in SELLER_PHONES:
        return f"+91{digits}"
    if len(digits) == 11 and digits.startswith("0"):  # landline
        return digits
    return ""  # reject seller numbers / invalid


def build_lead(idx: int, card: dict, panel: dict, modal: dict) -> dict:
    # Phone — prefer modal (more reliable), fall back to panel regex
    phone = clean_phone(modal.get("phone", "")) or clean_phone(panel.get("phone", ""))

    # Buyer name — prefer modal heading (most accurate)
    buyer_name = (modal.get("buyer_name") or panel.get("buyer_name_chat") or
                  card.get("name", ""))

    # Location from modal
    location = modal.get("location", "") or modal.get("address", "") or card.get("location", "")
    city = state = ""
    if location:
        parts = [p.strip() for p in location.split(",")]
        city  = parts[0] if len(parts) > 0 else ""
        state = parts[1] if len(parts) > 1 else ""
        state = re.sub(r"\d{6}", "", state).strip()

    # Messages
    msgs = panel.get("messages", [])
    message = "\n".join(msgs)[:2000]

    return {
        "_idx":            idx,
        "buyer_name":      buyer_name,
        "phone":           phone,
        "email":           modal.get("email", ""),
        "city":            city,
        "state":           state,
        "address":         location,
        "whatsapp_active": modal.get("whatsapp", ""),
        "product":         panel.get("product", ""),
        "quantity":        panel.get("quantity", ""),
        "grade":           panel.get("grade", ""),
        "buyer_filled":    panel.get("buyer_filled", ""),
        "order_value":     panel.get("order_value", ""),
        "req_type":        panel.get("req_type", ""),
        "packaging":       panel.get("packaging", ""),
        "message":         message,
        "req_count":       modal.get("req_count", ""),
        "calls_count":     modal.get("calls_count", ""),
        "replies_count":   modal.get("replies_count", ""),
        "platform_period": modal.get("platform_period", ""),
        "source":          panel.get("source", ""),
        "enquiry_date":    card.get("time", "") or panel.get("enquiry_date", ""),
        "scraped_at":      datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"),
    }


def save_xlsx(leads: list) -> Path:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    HEADERS = ["Buyer Name", "Phone", "Email", "City", "State", "Address",
               "WhatsApp Active", "Product", "Quantity", "Grade",
               "Buyer Filled Details", "Order Value", "Req Type", "Packaging",
               "Message", "Requirements (Engmt)", "Calls (Engmt)", "Replies (Engmt)",
               "Platform Period", "Source", "Enquiry Date", "Scraped At"]
    KEYS    = ["buyer_name", "phone", "email", "city", "state", "address",
               "whatsapp_active", "product", "quantity", "grade",
               "buyer_filled", "order_value", "req_type", "packaging",
               "message", "req_count", "calls_count", "replies_count",
               "platform_period", "source", "enquiry_date", "scraped_at"]

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "IndiaMART Leads"

    hfill = PatternFill("solid", fgColor="1F4E79")
    hfont = Font(color="FFFFFF", bold=True)
    for ci, h in enumerate(HEADERS, 1):
        c = ws.cell(1, ci, h)
        c.fill = hfill; c.font = hfont
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    ws.row_dimensions[1].height = 22

    df = Font(size=10)
    for ri, lead in enumerate(leads, 2):
        for ci, k in enumerate(KEYS, 1):
            v = str(lead.get(k) or "")
            cell = ws.cell(ri, ci, v); cell.font = df
            if k == "phone": cell.number_format = "@"
        ws.row_dimensions[ri].height = 28
        # wrap message column (O = col 15)
        ws.cell(ri, 15).alignment = Alignment(wrap_text=True, vertical="top")

    # Column widths
    col_widths = [28, 16, 28, 16, 16, 30, 22, 30, 12, 12,
                  30, 18, 18, 14, 50, 12, 12, 12, 18, 28, 14, 22]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Total Leads";   ws2["B1"] = len(leads)
    ws2["A2"] = "With Phone";    ws2["B2"] = sum(1 for l in leads if l.get("phone","").strip())
    ws2["A3"] = "With Email";    ws2["B3"] = sum(1 for l in leads if l.get("email","").strip())
    ws2["A4"] = "WhatsApp Active"; ws2["B4"] = sum(1 for l in leads if l.get("whatsapp_active","").strip())
    ws2["A5"] = "Scraped At";    ws2["B5"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")

    try:
        wb.save(OUT_XLSX)
        return OUT_XLSX
    except PermissionError:
        ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        fb = OUT_XLSX.with_name(f"indiamart_leads_{ts}.xlsx")
        wb.save(fb); print(f"  [!] File locked — saved as {fb.name}"); return fb


def save_csv(leads: list):
    FIELDS = ["buyer_name","phone","email","city","state","address",
              "whatsapp_active","product","quantity","grade",
              "buyer_filled","order_value","req_type","packaging",
              "message","req_count","calls_count","replies_count",
              "platform_period","source","enquiry_date","scraped_at"]
    try:
        with OUT_CSV.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
            w.writeheader(); w.writerows(leads)
    except PermissionError:
        ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        fb = OUT_CSV.with_name(f"indiamart_leads_{ts}.csv")
        with fb.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
            w.writeheader(); w.writerows(leads)
        print(f"  [!] CSV locked — saved as {fb.name}")


def auto_save(leads, found_phone, found_email):
    save_ckpt(leads); save_csv(leads); save_xlsx(leads)
    print(f"    >> Saved {len(leads)} leads | {found_phone} phone | {found_email} email")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    print("=" * 65)
    print("  IndiaMART Lead Manager Scraper — Contacts tab")
    print("=" * 65)

    all_leads    = load_ckpt()
    done_idxs    = {l["_idx"] for l in all_leads}
    found_phone  = sum(1 for l in all_leads if l.get("phone","").strip())
    found_email  = sum(1 for l in all_leads if l.get("email","").strip())
    print(f"  Checkpoint: {len(all_leads)} leads already saved")

    with sync_playwright() as pw:
        # ── Try to connect to your already-open Chrome via remote debug port ──
        # If Chrome isn't open with that port, fall back to launching fresh.
        DEBUG_PORT = 9222
        ctx = None
        page = None

        try:
            browser = pw.chromium.connect_over_cdp(f"http://localhost:{DEBUG_PORT}")
            ctx_obj = browser.contexts[0] if browser.contexts else None
            page = ctx_obj.pages[0] if (ctx_obj and ctx_obj.pages) else None
            if page:
                print(f"\n  Connected to your open Chrome! URL: {page.url[:70]}")
                USE_CDP = True
            else:
                raise Exception("No pages found")
        except Exception as e:
            print(f"\n  Could not connect to open Chrome ({e})")
            print("  Launching fresh Chrome with saved profile...")
            USE_CDP = False
            ctx = pw.chromium.launch_persistent_context(
                user_data_dir=str(PROFILE_DIR),
                headless=False,
                channel="chrome",
                args=["--start-maximized",
                      "--disable-blink-features=AutomationControlled",
                      "--disable-infobars",
                      "--no-first-run",
                      "--no-default-browser-check"],
                ignore_default_args=["--enable-automation"],
                no_viewport=True,
            )
            page = ctx.pages[0] if ctx.pages else ctx.new_page()

        page.set_default_timeout(60_000)
        page.add_init_script(
            "Object.defineProperty(navigator,'webdriver',{get:()=>undefined});"
        )

        # ── Check current page — if already on messagecentre with filter set, don't navigate ──
        time.sleep(4)
        cur_url = page.url.lower()
        print(f"\n  Current URL: {page.url[:80]}")

        if "messagecentre" in cur_url:
            print("  Already on Lead Manager — keeping your filter as-is")
            time.sleep(2)
        else:
            # Navigate fresh
            print(f"  Navigating to: {TARGET_URL}")
            try:
                page.goto(TARGET_URL, wait_until="domcontentloaded", timeout=60_000)
            except Exception:
                pass
            time.sleep(5)

            # Login wait (up to 3 min)
            if "messagecentre" not in page.url.lower():
                print("\n  LOGIN REQUIRED — log in to IndiaMART in Chrome (3 min)")
                for _ in range(36):
                    time.sleep(5)
                    if "messagecentre" in page.url.lower():
                        break
                    print(f"    waiting... {page.url[:60]}")
                else:
                    print("  Timeout."); ctx.close(); return
                time.sleep(2)

        print(f"  URL: {page.url[:80]}")

        # Click "Contacts" tab (safe — won't reset filter)
        try:
            ct = page.locator('[class*="cntmsg_tab"]:has-text("Contacts"), '
                              'button:has-text("Contacts"), '
                              'a:has-text("Contacts")').first
            if ct.count() > 0:
                ct.click(); time.sleep(2)
                print("  Contacts tab selected")
        except Exception:
            pass

        # Wait for list to render with the filter applied
        time.sleep(3)

        # Detect ReactVirtualized list
        time.sleep(2)
        info = page.evaluate(JS_LIST_INFO)
        if not info.get("ok"):
            print("  ReactVirtualized list not found!")
            if not USE_CDP and ctx: ctx.close()
            return

        row_h = info["rowH"]

        # ── Scroll to bottom repeatedly until no new items load ─────────────
        # The list lazy-loads contacts as you scroll — keep going until stable
        print(f"\n  Scrolling to load all contacts (rowHeight: {row_h}px)...")
        prev_total = 0
        stable_count = 0
        scroll_step = 0

        while stable_count < 6:   # stop after 6 consecutive scrolls with no new items
            # Scroll to absolute bottom
            page.evaluate(r"""
                () => {
                    var list = document.querySelector('.ReactVirtualized__List');
                    if (list) list.scrollTop = list.scrollHeight + 99999;
                }
            """)
            time.sleep(1.2)   # reduced from 2.0 — faster scroll loading

            info = page.evaluate(JS_LIST_INFO)
            cur_total = info.get("total", 0)
            scroll_step += 1

            if cur_total > prev_total:
                print(f"  Step {scroll_step:3d}: {cur_total} contacts loaded (+{cur_total - prev_total})")
                prev_total = cur_total
                stable_count = 0   # reset — new items appeared
            else:
                stable_count += 1

        total = prev_total
        print(f"\n  All contacts loaded: {total} total")

        # Scroll back to top before scraping
        page.evaluate(r"""
            () => {
                var list = document.querySelector('.ReactVirtualized__List');
                if (list) list.scrollTop = 0;
            }
        """)
        time.sleep(2)
        print(f"  Scraping {total - len(done_idxs)} new leads\n")
        print(f"  {'#':>4}  {'Company':25}  {'Buyer':25}  {'Phone':16}  {'Product':25}")
        print(f"  {'-'*4}  {'-'*25}  {'-'*25}  {'-'*16}  {'-'*25}")
        for idx in range(total):
            if idx in done_idxs:
                continue

            # Scroll card into view
            page.evaluate(JS_SCROLL_TO, idx)
            time.sleep(1.2)

            # Get card info
            card = page.evaluate(JS_GET_CARD, idx)
            if card is None:
                # Try once more after extra wait
                time.sleep(1.5)
                card = page.evaluate(JS_GET_CARD, idx)
            if card is None:
                print(f"  {idx+1:>4}  skip — card not rendered")
                continue

            company = (card.get("name") or "?")[:25]

            # Click card
            clicked = page.evaluate(JS_CLICK_CARD, idx)
            if not clicked:
                print(f"  {idx+1:>4}  {company:25}  click failed")
                continue
            time.sleep(2)

            # Scrape right panel
            panel = page.evaluate(JS_SCRAPE_PANEL)

            # Click "View More"
            vm_result = page.evaluate(JS_CLICK_VIEWMORE)
            modal = {}
            if vm_result.get("clicked"):
                time.sleep(2.5)
                modal = page.evaluate(JS_SCRAPE_MODAL)
                if not modal.get("ok"):
                    # Print diagnostic on first failure
                    diag = modal.get("diag", [])
                    if diag:
                        print(f"\n    [!] Modal not found. High-z elements:")
                        for d in diag[:3]:
                            print(f"        z={d['z']} {d['w']}x{d['h']} {d['cls'][:70]}")

                # Close modal
                closed = page.evaluate(JS_CLOSE_MODAL)
                if not closed.get("closed"):
                    page.keyboard.press("Escape")
                time.sleep(0.8)

            # Build record
            lead = build_lead(idx, card, panel, modal)

            buyer   = (lead["buyer_name"] or "?")[:25]
            phone   = lead["phone"] or "-"
            product = (lead["product"] or "?")[:25]
            vm_ok   = "✓" if modal.get("ok") else "–"
            print(f"  {idx+1:>4}  {company:25}  {buyer:25}  {phone:16}  {product}  vm={vm_ok}")

            if lead["phone"]: found_phone += 1
            if lead["email"]: found_email += 1

            all_leads.append(lead)
            done_idxs.add(idx)

            if len(all_leads) % 5 == 0:
                auto_save(all_leads, found_phone, found_email)

            time.sleep(1.0 + random.uniform(0, 0.6))

        # Final save
        auto_save(all_leads, found_phone, found_email)
        out = save_xlsx(all_leads)
        # Only close if we launched our own Chrome; don't close the user's browser
        if not USE_CDP and ctx:
            ctx.close()

    print(f"\n{'='*65}")
    print(f"  DONE")
    print(f"  Total leads : {len(all_leads)}")
    print(f"  With phone  : {found_phone}")
    print(f"  With email  : {found_email}")
    print(f"  Excel       : {out}")
    print(f"  CSV         : {OUT_CSV}")
    print(f"{'='*65}\n")


if __name__ == "__main__":
    main()
