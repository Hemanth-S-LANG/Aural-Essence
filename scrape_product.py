"""
Generic TradeIndia product scraper.
Usage: python scrape_product.py "Rosemary Hair Oil"

Steps:
  1. Scroll through /manufacturers/{slug}.html until Show More is gone.
  2. Extract company names + profile URLs from DOM.
  3. Fetch each profile page with requests to get contact person, address,
     city, state, GST no.
  4. Save to output/{slug}_tradeindia.xlsx

Collects (all publicly visible without login):
  Company Name, Contact Person, Designation, Address, City, State,
  GST No, Product Name, Product URL, Profile URL, Phone Status, Scraped At

Phone is always HIDDEN on TradeIndia (number_mask=true for all sellers).
Email is always MISSING (never shown publicly).
Brand Website / LinkedIn / Marketplace: not present on TradeIndia pages.
"""
import csv, json, re, sys, time, random
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

# ── Config ────────────────────────────────────────────────────────────────────
PRODUCT    = sys.argv[1] if len(sys.argv) > 1 else "Rosemary Hair Oil"
SLUG       = re.sub(r"\s+", "-", PRODUCT.strip().lower())
BASE_URL   = "https://www.tradeindia.com"
LIST_URL   = f"{BASE_URL}/manufacturers/{SLUG}.html"
UA_BROWSER = "AuralEssenceScraper/1.0 (+research-bot; contact: intern@auralEssence.com)"
UA_REQ     = "AuralEssenceScraper/1.0 (+research-bot; contact: intern@auralEssence.com)"
DELAY_MIN  = 3.5
DELAY_MAX  = 5.0

OUT_DIR    = Path("output"); OUT_DIR.mkdir(exist_ok=True)
CKPT_DIR   = Path("data");   CKPT_DIR.mkdir(exist_ok=True)
TS         = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
OUT_CSV    = OUT_DIR / f"{SLUG}_tradeindia_{TS}.csv"
OUT_XLSX   = OUT_DIR / f"{SLUG}_tradeindia.xlsx"
CKPT_FILE  = CKPT_DIR / f"{SLUG}_profile_checkpoint.json"

FIELDS = ["company_name","contact_person","designation","address",
          "city","state","gst_no","product_name","product_url",
          "profile_url","phone_status","email_status","scraped_at"]

print("=" * 65)
print(f"TradeIndia scraper — {PRODUCT}")
print(f"Listing URL : {LIST_URL}")
print(f"Output      : {OUT_XLSX}")
print("=" * 65)

# ── JS helpers ────────────────────────────────────────────────────────────────
_JS_EXTRACT = r"""
() => {
    const BASE = 'https://www.tradeindia.com';
    const results = [];
    const seenProfiles = new Set();

    document.querySelectorAll('a[href]').forEach(a => {
        const href = a.href || '';
        if (!/tradeindia\.com\/[a-z0-9\-]+-\d+\/$/.test(href)) return;
        if (seenProfiles.has(href)) return;
        seenProfiles.add(href);

        let card = a.parentElement;
        for (let i = 0; i < 10 && card; i++) {
            const cls = (card.className||'').toLowerCase();
            const tag = card.tagName;
            if ((tag==='LI'||tag==='DIV') &&
                (cls.includes('mcat')||cls.includes('supplier')||
                 cls.includes('listing')||cls.includes('card')||
                 cls.includes('product'))) break;
            card = card.parentElement;
        }

        const coName = a.innerText.trim() || a.title || '';
        let productName='', productUrl='';
        if (card) {
            const pl = card.querySelector('a[href*="/products/"]');
            if (pl) { productUrl=pl.href; productName=(pl.innerText||pl.title||'').trim(); }
        }

        let city='', state='';
        if (card) {
            const cardText = card.innerText||'';
            const m = cardText.match(/\|\s*([A-Za-z\s]+),\s*([A-Za-z\s]+),\s*India/);
            if (m) { city=m[1].trim(); state=m[2].trim(); }
        }

        results.push({ profile_url:href, company_name:coName,
                       product_name:productName, product_url:productUrl,
                       city:city, state:state });
    });

    let total_pages=0;
    try {
        const el=document.getElementById('__NEXT_DATA__');
        if(el){ const nd=JSON.parse(el.textContent);
            const pag=nd?.props?.pageProps?.initialState
                         ?.manufacturer?.manufacturer_page
                         ?.manufacturer_page_res?.mcatData
                         ?.listing_data?.pagination;
            if(pag?.total_pages) total_pages=pag.total_pages; }
    } catch(e){}

    const quoteCount=(document.body.innerText.match(/Get Best Quote/g)||[]).length;
    const showMoreVisible=!!document.querySelector('button.more-listing-btn');

    return JSON.stringify({ entries:results, quote_count:quoteCount,
                            total_pages:total_pages, show_more_visible:showMoreVisible });
}
"""

_JS_WAIT_QUOTES = """
(n) => new Promise(resolve => {
    const deadline = Date.now() + 25000;
    const check = () => {
        const c=(document.body.innerText.match(/Get Best Quote/g)||[]).length;
        if(c>n){resolve(c);return;}
        if(Date.now()>deadline){resolve(-1);return;}
        setTimeout(check,600);
    };
    setTimeout(check,600);
})
"""

# ── Phase 1: Scroll collection ────────────────────────────────────────────────
def phase1_collect():
    all_entries = {}
    prev_quotes = 0
    consec_empty = 0
    total_pages  = 0

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=False)
        ctx = browser.new_context(user_agent=UA_BROWSER, locale="en-IN",
                                  extra_http_headers={"Accept-Language":"en-IN,en;q=0.9"})
        page = ctx.new_page()
        page.set_default_timeout(60_000)

        print(f"\nPhase 1: Loading {LIST_URL} ...")
        loaded = False
        for attempt in range(1, 6):
            try:
                page.goto(LIST_URL, wait_until="domcontentloaded", timeout=60_000)
                time.sleep(5)
                print(f"  Loaded (attempt {attempt}): {page.title()!r}")
                loaded = True; break
            except Exception as e:
                wait = attempt * 15
                print(f"  Attempt {attempt} failed. Waiting {wait}s..."); time.sleep(wait)

        if not loaded:
            print("  All attempts failed."); ctx.close(); browser.close(); return {}

        # Initial extraction
        raw0 = page.evaluate(_JS_EXTRACT)
        d0   = json.loads(raw0)
        total_pages = d0.get("total_pages", 0)
        prev_quotes = d0.get("quote_count", 0)
        ts0         = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")

        for e in d0.get("entries", []):
            key = e["profile_url"] or e["company_name"]
            if key: all_entries[key] = {**e, "phone_status":"HIDDEN",
                                         "email_status":"MISSING","scraped_at":ts0}

        print(f"  Initial: {len(all_entries)} entries | quotes={prev_quotes} | "
              f"total_pages={total_pages} | show_more={d0.get('show_more_visible')}")

        if total_pages == 0 and len(all_entries) == 0:
            print(f"\n  NOTE: No listings found for '{PRODUCT}' on TradeIndia.")
            print(f"  URL tried: {LIST_URL}")
            print("  The product may be listed under a different category name.")
            ctx.close(); browser.close(); return {}

        # Click Show More loop
        click_num = 0
        while click_num < 100:
            click_num += 1

            btn = page.locator("button.more-listing-btn")
            if btn.count() == 0:
                page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                time.sleep(3)
                if btn.count() == 0:
                    consec_empty += 1
                    if consec_empty >= 5:
                        print(f"  No Show More button — all content loaded.")
                        break
                    continue
            else:
                consec_empty = 0

            time.sleep(DELAY_MIN + random.uniform(0, DELAY_MAX - DELAY_MIN))

            # Dismiss any popup overlay
            page.keyboard.press("Escape")
            time.sleep(0.5)
            try:
                for sel in ["[class*='popClose']","[class*='crossBtn']",
                            "button[class*='close']","[aria-label='Close']"]:
                    if page.locator(sel).count() > 0:
                        page.locator(sel).first.click(timeout=2000)
                        time.sleep(0.3); break
            except Exception: pass
            try:
                page.mouse.click(10, 10); time.sleep(0.5)
            except Exception: pass

            # Click via dispatch_event (bypasses overlay pointer interception)
            try:
                b = page.locator("button.more-listing-btn")
                if b.count() > 0:
                    b.first.scroll_into_view_if_needed()
                    time.sleep(0.3)
                    b.first.dispatch_event("click")
            except Exception:
                page.evaluate("window.scrollTo(0, document.body.scrollHeight)")

            new_q = page.evaluate(_JS_WAIT_QUOTES, prev_quotes)
            try: page.wait_for_load_state("networkidle", timeout=8_000)
            except PWTimeout: pass

            raw2 = page.evaluate(_JS_EXTRACT)
            d2   = json.loads(raw2)
            entries2   = d2.get("entries", [])
            quotes2    = d2.get("quote_count", 0)
            show_more2 = d2.get("show_more_visible", False)
            if d2.get("total_pages"): total_pages = d2["total_pages"]

            added = 0; ts2 = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
            for e in entries2:
                key = e["profile_url"] or e["company_name"]
                if key and key not in all_entries:
                    all_entries[key] = {**e, "phone_status":"HIDDEN",
                                        "email_status":"MISSING","scraped_at":ts2}
                    added += 1

            prev_quotes = quotes2
            print(f"  Click {click_num:3d}: +{added:3d} new | quotes={quotes2:4d} | "
                  f"total unique={len(all_entries):4d} | show_more={show_more2}")

            if not show_more2 and added == 0:
                print("  Show More gone — all pages loaded!"); break

        ctx.close(); browser.close()

    print(f"\nPhase 1 complete: {len(all_entries)} unique sellers collected")
    return all_entries

# ── Phase 2: Fetch profiles ───────────────────────────────────────────────────
def phase2_enrich(all_entries: dict) -> dict:
    # Load existing checkpoint
    checkpoint = {}
    if CKPT_FILE.exists():
        try:
            checkpoint = json.loads(CKPT_FILE.read_text(encoding="utf-8"))
            print(f"Checkpoint: {len(checkpoint)} profiles already fetched")
        except Exception: pass

    sess = requests.Session()
    sess.headers.update({"User-Agent": UA_REQ,
                         "Accept": "text/html,application/xhtml+xml;q=0.9,*/*;q=0.8",
                         "Accept-Language": "en-IN,en;q=0.9",
                         "Accept-Encoding": "gzip, deflate, br"})

    todo = [url for url in all_entries
            if url.startswith("http") and url.rstrip("/") not in checkpoint]
    print(f"\nPhase 2: Fetching {len(todo)} profile pages...")

    fetched = errors = consec_err = 0
    for i, url in enumerate(todo, 1):
        key = url.rstrip("/")
        delay = DELAY_MIN + random.uniform(0, DELAY_MAX - DELAY_MIN)
        time.sleep(delay)

        if i % 25 == 1:
            print(f"  [{len(checkpoint)+fetched:4d}/{len(all_entries)}] {url[:65]}")

        try:
            r = sess.get(url, timeout=20, allow_redirects=True)
            if r.status_code != 200:
                checkpoint[key] = {"_status": r.status_code}
                errors += 1; consec_err += 1
                if consec_err >= 5:
                    print(f"  5 consecutive errors — pausing 30s"); time.sleep(30); consec_err = 0
                continue

            soup   = BeautifulSoup(r.text, "lxml")
            nd_tag = soup.find("script", id="__NEXT_DATA__")
            if not nd_tag:
                checkpoint[key] = {"_status": "no_next_data"}; errors += 1; continue

            nd    = json.loads(nd_tag.string)
            state = (nd.get("props",{}).get("pageProps",{}).get("initialState",{}))
            cd    = (state.get("sellerProfile",{})
                         .get("seller_profile",{})
                         .get("seller_profile_res",{})
                         .get("company_details_data",{}))
            if not cd:
                checkpoint[key] = {"_status": "no_company_data"}; errors += 1; continue

            # Get product from sellerMoreProducts
            more = (state.get("sellerMoreProducts",{})
                         .get("seller_more_products",{})
                         .get("seller_more_products_res",{})
                         .get("data", []))
            prod_name = cd.get("product_name","")
            prod_url  = ""
            if more and isinstance(more, list) and more[0]:
                p0        = more[0]
                prod_name = prod_name or p0.get("product_name","")
                pp        = p0.get("prod_url","")
                prod_url  = (BASE_URL + pp if pp.startswith("/") else pp) if pp else ""

            checkpoint[key] = {
                "_status":        "ok",
                "contact_person": cd.get("owner_name",""),
                "designation":    cd.get("desg",""),
                "address":        cd.get("address",""),
                "city":           cd.get("city",""),
                "state":          cd.get("state",""),
                "gst_no":         cd.get("gst_no",""),
                "product_name":   prod_name,
                "product_url":    prod_url,
            }
            fetched += 1; consec_err = 0

            if i % 10 == 0:
                print(f"  [{len(checkpoint):4d}/{len(all_entries)}] "
                      f"fetched={fetched}, errors={errors}")

        except Exception as e:
            checkpoint[key] = {"_status": f"error:{e}"}
            errors += 1; consec_err += 1

        if fetched % 25 == 0 and fetched > 0:
            CKPT_FILE.write_text(json.dumps(checkpoint), encoding="utf-8")

    CKPT_FILE.write_text(json.dumps(checkpoint), encoding="utf-8")
    print(f"Phase 2 complete: {fetched} OK, {errors} errors")
    return checkpoint

# ── Phase 3: Merge + Export ───────────────────────────────────────────────────
def phase3_export(all_entries: dict, checkpoint: dict):
    rows = list(all_entries.values())

    # Merge enrichment
    for row in rows:
        url  = row.get("profile_url","").rstrip("/")
        data = checkpoint.get(url) or checkpoint.get(url+"/")
        if data and data.get("_status") == "ok":
            for f in ["contact_person","designation","address",
                      "city","state","gst_no","product_name","product_url"]:
                if not row.get(f):
                    row[f] = data.get(f,"")

    # Stats
    print(f"\nData quality:")
    for f in ["company_name","contact_person","city","state",
              "address","gst_no","product_name"]:
        n = sum(1 for r in rows if str(r.get(f,"") or "").strip())
        print(f"  {f:<20}: {n}/{len(rows)}")

    # CSV
    with OUT_CSV.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)
    print(f"\nCSV saved: {OUT_CSV}")

    # Excel
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = f"{PRODUCT} Sellers"

    headers = ["Company Name","Contact Person","Designation",
               "Address","City","State","GST No",
               "Product Name","Product URL","Profile URL",
               "Phone Status","Email Status","Scraped At"]
    col_keys = ["company_name","contact_person","designation",
                "address","city","state","gst_no",
                "product_name","product_url","profile_url",
                "phone_status","email_status","scraped_at"]

    fill = PatternFill("solid", fgColor="1F4E79")
    hf   = Font(color="FFFFFF", bold=True)
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h); c.fill=fill; c.font=hf
        c.alignment=Alignment(horizontal="center")
    ws.freeze_panes="A2"
    ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}1"
    lf=Font(color="0563C1",underline="single",size=10)
    df=Font(size=10)

    for ri,row in enumerate(rows,2):
        for ci,key in enumerate(col_keys,1):
            v=str(row.get(key,"") or "")
            cell=ws.cell(ri,ci,v); cell.font=df
            if key in ("product_url","profile_url") and v:
                cell.font=lf; cell.hyperlink=v

    for i,key in enumerate(col_keys,1):
        col_vals=[str(ws.cell(r,i).value or "") for r in range(2,ws.max_row+1)]
        w2=min(max(len(headers[i-1])+2,
                   max((len(v) for v in col_vals if v),default=0)+2),60)
        ws.column_dimensions[get_column_letter(i)].width=w2

    ws2=wb.create_sheet("Run Summary")
    ws2["A1"]="Product";         ws2["B1"]=PRODUCT
    ws2["A2"]="Source";          ws2["B2"]="TradeIndia"
    ws2["A3"]="Last Updated";    ws2["B3"]=datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
    ws2["A4"]="Total Sellers";   ws2["B4"]=len(rows)
    ws2["A5"]="Unique Companies";ws2["B5"]=len(set(r.get("company_name","") for r in rows))
    ws2["A6"]="City filled";     ws2["B6"]=sum(1 for r in rows if str(r.get("city","") or "").strip())
    ws2["A7"]="Contact filled";  ws2["B7"]=sum(1 for r in rows if str(r.get("contact_person","") or "").strip())
    ws2["A8"]="Phone Status";    ws2["B8"]="All HIDDEN (masked by TradeIndia)"
    ws2["A9"]="Email Status";    ws2["B9"]="All MISSING (not shown on TradeIndia)"

    try:
        wb.save(OUT_XLSX)
    except PermissionError:
        ts2=datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        OUT_XLSX_FALLBACK=OUT_DIR/f"{SLUG}_tradeindia_{ts2}.xlsx"
        wb.save(OUT_XLSX_FALLBACK)
        print(f"Excel locked — saved to: {OUT_XLSX_FALLBACK}")
        return

    print(f"Excel saved: {OUT_XLSX}  ({len(rows)} rows)")
    print(f"\nFirst 10 rows:")
    print(f"{'#':>4}  {'Company':<38}  {'Product':<30}  {'City':<14}  Contact")
    print("-"*105)
    for i,r in enumerate(rows[:10],1):
        print(f"{i:>4}  {str(r.get('company_name',''))[:37]:<38}  "
              f"{str(r.get('product_name',''))[:29]:<30}  "
              f"{str(r.get('city',''))[:13]:<14}  "
              f"{str(r.get('contact_person',''))[:25]}")

# ─────────────────────────────────────────────────────────────────────────────
all_entries = phase1_collect()
if all_entries:
    checkpoint = phase2_enrich(all_entries)
    phase3_export(all_entries, checkpoint)
else:
    print("\nNo data collected — check product name or try a different slug.")
    print(f"URL attempted: {LIST_URL}")
