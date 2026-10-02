"""
Tests for:
  - --source flag (CLI argument parsing)
  - Per-source scraper selection
  - Per-source Excel file output
  - Meesho blocked sheet output
  - CSV import command
"""
import csv
import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock
from datetime import datetime

import openpyxl

from app.main import _parse_args, _select_scrapers, run
from app.scrapers.tradeindia_scraper import TradeIndiaScraper
from app.scrapers.meesho_scraper import MeeshoScraper
from app.exporters.excel_exporter import export_meesho_blocked
from app.importers.csv_importer import import_csv
from app.models.lead import Lead


# ── CLI argument parsing ──────────────────────────────────────────────────────

class TestArgParsing:
    def test_default_source_is_all(self):
        args = _parse_args([])
        assert args.source == "all"

    def test_source_tradeindia(self):
        args = _parse_args(["--source", "tradeindia"])
        assert args.source == "tradeindia"

    def test_source_meesho(self):
        args = _parse_args(["--source", "meesho"])
        assert args.source == "meesho"

    def test_source_all_explicit(self):
        args = _parse_args(["--source", "all"])
        assert args.source == "all"

    def test_invalid_source_raises(self):
        with pytest.raises(SystemExit):
            _parse_args(["--source", "invalid"])

    def test_import_csv_arg(self):
        args = _parse_args(["--import-csv", "data.csv", "--source", "meesho"])
        assert args.import_csv == "data.csv"
        assert args.source == "meesho"

    def test_import_csv_no_source_defaults_all(self):
        args = _parse_args(["--import-csv", "data.csv"])
        assert args.source == "all"


# ── Scraper selection ─────────────────────────────────────────────────────────

class TestScraperSelection:
    def test_all_returns_both(self):
        scrapers = _select_scrapers("all")
        names = {s.source_name for s in scrapers}
        assert "TradeIndia" in names
        assert "Meesho" in names

    def test_tradeindia_returns_one(self):
        scrapers = _select_scrapers("tradeindia")
        assert len(scrapers) == 1
        assert scrapers[0].source_name == "TradeIndia"

    def test_meesho_returns_one(self):
        scrapers = _select_scrapers("meesho")
        assert len(scrapers) == 1
        assert scrapers[0].source_name == "Meesho"


# ── Per-source Excel output ───────────────────────────────────────────────────

class TestPerSourceExport:
    def _ti_lead(self, name="Earth Science Ayurveda"):
        return Lead(
            company_name=name,
            source="TradeIndia",
            product_name="Rose Water",
            product_url="https://www.tradeindia.com/products/rose-water.html",
            phone_status="HIDDEN",
            email_status="MISSING",
        )

    def test_tradeindia_file_created(self, tmp_path):
        from app.exporters.excel_exporter import export_excel
        out = tmp_path / "ti.xlsx"
        export_excel([self._ti_lead()], output_path=out, source_label="tradeindia")
        assert out.exists()

    def test_tradeindia_file_has_data(self, tmp_path):
        from app.exporters.excel_exporter import export_excel
        out = tmp_path / "ti.xlsx"
        export_excel([self._ti_lead()], output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        assert ws.max_row == 2  # header + 1 data row

    def test_meesho_blocked_sheet_has_two_sheets(self, tmp_path):
        out = tmp_path / "meesho_blocked.xlsx"
        export_meesho_blocked(output_path=out)
        wb = openpyxl.load_workbook(out)
        assert "Leads" in wb.sheetnames
        assert "Status" in wb.sheetnames

    def test_meesho_blocked_leads_sheet_header_only(self, tmp_path):
        out = tmp_path / "meesho_blocked.xlsx"
        export_meesho_blocked(output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Leads"]
        assert ws.max_row == 1   # header only, no data rows

    def test_meesho_blocked_status_sheet_contains_reason(self, tmp_path):
        out = tmp_path / "meesho_blocked.xlsx"
        export_meesho_blocked(output_path=out)
        wb = openpyxl.load_workbook(out)
        ws2 = wb["Status"]
        all_text = " ".join(
            str(ws2.cell(r, c).value or "")
            for r in range(1, ws2.max_row + 1)
            for c in (1, 2)
        )
        assert "BLOCKED" in all_text
        assert "WAF" in all_text or "Akamai" in all_text or "403" in all_text

    def test_meesho_status_mentions_manual_import(self, tmp_path):
        out = tmp_path / "meesho_blocked.xlsx"
        export_meesho_blocked(output_path=out)
        wb = openpyxl.load_workbook(out)
        ws2 = wb["Status"]
        all_text = " ".join(
            str(ws2.cell(r, c).value or "")
            for r in range(1, ws2.max_row + 1)
            for c in (1, 2)
        )
        assert "import-csv" in all_text or "import_csv" in all_text or "manual" in all_text.lower()


# ── CSV import ────────────────────────────────────────────────────────────────

def _write_csv(path: Path, rows: list, headers: list) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        w.writerows(rows)


class TestCsvImport:
    def test_valid_rows_imported(self, tmp_path):
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "Meena Perfumery", "phone_number": "9876543210",
             "email": "contact@gmail.com", "product_name": "Rose Water",
             "address": "Kannauj, UP"},
        ], ["company_name", "phone_number", "email", "product_name", "address"])
        leads, imported, skipped = import_csv(csv_path, "Meesho")
        assert imported == 1
        assert skipped == 0
        assert leads[0].company_name == "Meena Perfumery"

    def test_phone_validated(self, tmp_path):
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "Test Co", "phone_number": "9876543210"},
        ], ["company_name", "phone_number"])
        leads, _, _ = import_csv(csv_path, "Meesho")
        assert leads[0].phone_status == "VALID"
        assert leads[0].phone_number == "+919876543210"

    def test_phone_source_is_manual(self, tmp_path):
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "Test Co", "phone_number": "9876543210"},
        ], ["company_name", "phone_number"])
        leads, _, _ = import_csv(csv_path, "Meesho")
        assert leads[0].phone_source == "MANUAL"
        assert leads[0].email_source == "MANUAL"

    def test_source_set_correctly(self, tmp_path):
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "Test Co"},
        ], ["company_name"])
        leads, _, _ = import_csv(csv_path, "Meesho")
        assert leads[0].source == "Meesho"

    def test_missing_company_skipped(self, tmp_path):
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "", "phone_number": "9876543210"},
            {"company_name": "Valid Co", "phone_number": ""},
        ], ["company_name", "phone_number"])
        leads, imported, skipped = import_csv(csv_path, "Meesho")
        assert imported == 1
        assert skipped == 1
        assert leads[0].company_name == "Valid Co"

    def test_invalid_email_stored_as_invalid(self, tmp_path):
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "Test Co", "email": "notanemail"},
        ], ["company_name", "email"])
        leads, _, _ = import_csv(csv_path, "Meesho")
        assert leads[0].email_status == "INVALID"
        assert leads[0].email is None

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            import_csv(tmp_path / "nonexistent.csv", "Meesho")

    def test_empty_csv_returns_zero(self, tmp_path):
        csv_path = tmp_path / "empty.csv"
        _write_csv(csv_path, [], ["company_name", "phone_number"])
        leads, imported, skipped = import_csv(csv_path, "Meesho")
        assert imported == 0
        assert skipped == 0
        assert leads == []

    def test_extra_columns_ignored(self, tmp_path):
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "Test Co", "unknown_col": "ignored", "random": "value"},
        ], ["company_name", "unknown_col", "random"])
        leads, imported, skipped = import_csv(csv_path, "Meesho")
        assert imported == 1
        assert leads[0].company_name == "Test Co"

    def test_hidden_phone_placeholder_stored_as_hidden(self, tmp_path):
        """CSV rows with 'View Number' must set HIDDEN, not attempt bypass."""
        csv_path = tmp_path / "test.csv"
        _write_csv(csv_path, [
            {"company_name": "Test Co", "phone_number": "View Number"},
        ], ["company_name", "phone_number"])
        leads, _, _ = import_csv(csv_path, "Meesho")
        assert leads[0].phone_status == "HIDDEN"
        assert leads[0].phone_number is None
